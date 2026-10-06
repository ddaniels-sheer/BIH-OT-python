import zipfile
import pandas as pd
from pathlib import Path
from openpyxl.styles import Alignment, PatternFill

zip_path = r"C:\Users\DiamonDaniels\test.zip"        # zipped folder
output_path = r"output.xlsx"                              # the excel file to create
#colors = ["#E99299D1", "#00FF00", "#92CEE0", "#FDFFA4FF", "#F1C4E2", "#cdb4e8"]  # list of colors for sheets

ORANGE = "F4B084"
LIGHT_ORANGE = "FFD589"
BLUE = "ADD3E6"
LIGHT_BLUE = "92CEE0"
YELLOW = "FFFFE0"
GREEN = "93BC8F"
PINK = "F1C4E2"
PURPLE = "CDB4E8"
GRAY = "D9D9D9"

SHEET_CONFIG = {
    "Summary": {
        "tab_color": PINK,
        "header_group":{1: [("B", "D", ORANGE), ("E", "G", BLUE), ("H", "J", YELLOW), ("K", "M", GREEN), ("N", "R", PURPLE), ("S", "U", ORANGE),
                         ("V", "X", BLUE), ("Y", "AA", YELLOW), ("AB", "AD", GREEN), ("AE", "AI", PURPLE)]},
    },
    "LRC by Week": {
        "tab_color": PURPLE,
        "header_group": {1: [("B", "K", ORANGE), ("L", "U", BLUE)],
                         2: [("B", "F", LIGHT_ORANGE), ("G", "K", GRAY), ("L", "P", LIGHT_BLUE), ("Q", "U", GRAY)]},
    },
    "Rochester": {
        "tab_color": ORANGE,
        "header_group": {1: [("AA", "AA", GREEN), ("AB", "AB", GREEN), ("AR", "AR", YELLOW), ("AS", "AS", YELLOW)]}
    },
    "Sauget": {
        "tab_color": BLUE,
        "header_group": {1: [("AA", "AA", GREEN), ("AB", "AB", GREEN),("AC", "AC", PINK), ("AR", "AR", YELLOW), ("AS", "AS", YELLOW)]}
    },
    "Valmeyer": {
        "tab_color": YELLOW,
        "header_group": {1: [("AA", "AA", GREEN), ("AB", "AB", GREEN), ("AC", "AC", PINK), ("AR", "AR", YELLOW), ("AS", "AS", YELLOW)]}
    },
    "Other": {
        "tab_color": GREEN,
        "header_group": {1: [("AA", "AA", GREEN), ("AB", "AB", GREEN), ("AC", "AC", PINK), ("AR", "AR", YELLOW), ("AS", "AS", YELLOW)]}
    }
}

def format_sheet(ws, config):
    ws.sheet_properties.tabColor = config["tab_color"]
    for row_num, header_groups in config["header_group"].items():
        for start_col, end_col, color in header_groups:
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            for row in ws[f"{start_col}{row_num}:{end_col}{row_num}"]:
                for cell in row:
                    cell.fill = fill
            if start_col != end_col:
                ws.merge_cells(f"{start_col}{row_num}:{end_col}{row_num}")
                ws[f"{start_col}{row_num}"].alignment = Alignment(horizontal="center")

with zipfile.ZipFile(zip_path) as zf, pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    #i = 0
    for name in zf.namelist():
        # Skip anything that isn't a CSV
        if name.endswith("/") or not name.lower().endswith(".csv"):
            continue

        with zf.open(name) as f:
            df = pd.read_csv(f)
        df.columns = ["" if str(c).startswith("Unnamed") else c for c in df.columns]    #turn auto-generated unnamed columns back into empty strings

        # Excel sheet names: max of 31 chars, no special characters
        sheet_name = Path(name).stem[:31]
        for ch in '[]:*?/\\':
            sheet_name = sheet_name.replace(ch, "_")

        df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]
        #ws.tab_color = colors[i % len(colors)]
        if sheet_name in SHEET_CONFIG:
            format_sheet(ws, SHEET_CONFIG[sheet_name])
      
        #i += 1      # move to the next color for the next sheet (might not be needed if we'll be creating each sheet's own function)
        print(f"Added sheet: {sheet_name} ({len(df)} rows)")